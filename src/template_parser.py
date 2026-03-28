"""PPT Template Parser for extracting slide layouts and placeholders."""

from pathlib import Path
from typing import Dict, List, Any, Optional
import logging

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
except ImportError as e:
    raise ImportError(f"python-pptx is required: {e}")

logger = logging.getLogger(__name__)


class TemplateParser:
    """Parses PowerPoint templates to extract layout and placeholder information."""

    def __init__(self, template_path: str | Path):
        """
        Initialize TemplateParser with a PPTX file.

        Args:
            template_path: Path to the PowerPoint template file

        Raises:
            FileNotFoundError: If template file doesn't exist
            ValueError: If template path is invalid
        """
        self.template_path = Path(template_path)
        if not self.template_path.exists():
            raise FileNotFoundError(f"Template not found: {self.template_path}")
        if not self.template_path.suffix.lower() == ".pptx":
            raise ValueError(f"File must be a .pptx: {self.template_path}")

        self.presentation: Optional[Presentation] = None
        self.metadata: Dict[str, Any] = {}

    def parse(self) -> Dict[str, Any]:
        """
        Parse the PPTX template and extract structured metadata.

        Returns:
            Dictionary containing template metadata:
            {
                'template_name': str,
                'template_path': str,
                'slide_layouts': [
                    {
                        'name': str,
                        'layout_id': int,
                        'placeholders': [
                            {
                                'type': str,
                                'idx': int,
                                'width': float,
                                'height': float,
                                'left': float,
                                'top': float
                            }
                        ]
                    }
                ]
            }

        Raises:
            RuntimeError: If parsing fails
        """
        try:
            self.presentation = Presentation(str(self.template_path))
            self._extract_metadata()
            return self.metadata
        except Exception as e:
            logger.error(f"Failed to parse template {self.template_path}: {e}")
            raise RuntimeError(f"Template parsing failed: {e}") from e

    def _extract_metadata(self) -> None:
        """Extract metadata from presentation."""
        self.metadata = {
            "template_name": self.template_path.name,
            "template_path": str(self.template_path),
            "slide_layouts": [],
        }

        for layout_idx, layout in enumerate(self.presentation.slide_layouts):
            layout_info = {
                "name": layout.name,
                "layout_id": layout_idx,
                "placeholders": self._extract_placeholders(layout),
            }
            self.metadata["slide_layouts"].append(layout_info)

    def _extract_placeholders(self, layout) -> List[Dict[str, Any]]:
        """
        Extract placeholder information from a slide layout.

        Args:
            layout: A slide layout from python-pptx

        Returns:
            List of placeholder dictionaries with type and geometry
        """
        placeholders = []

        for shape in layout.shapes:
            if not shape.is_placeholder:
                continue

            phf = shape.placeholder_format
            try:
                placeholder_type = self._get_placeholder_type(phf.type)
                # Access idx which may raise KeyError if placeholder index is invalid
                idx = phf.idx
            except (ValueError, KeyError):
                placeholder_type = "UNKNOWN"
                try:
                    idx = phf.idx
                except (ValueError, KeyError):
                    # Skip this placeholder if idx is also inaccessible
                    continue

            placeholder_info = {
                "type": placeholder_type,
                "idx": idx,
                "width": shape.width.pt if shape.width else None,
                "height": shape.height.pt if shape.height else None,
                "left": shape.left.pt if shape.left else None,
                "top": shape.top.pt if shape.top else None,
            }
            placeholders.append(placeholder_info)

        return placeholders

    def _get_placeholder_type(self, ph_type: int) -> str:
        """
        Convert python-pptx placeholder type enum to string.

        Args:
            ph_type: Placeholder type enum from PP_PLACEHOLDER

        Returns:
            String representation of placeholder type
        """
        # Build mapping dynamically to handle version differences in PP_PLACEHOLDER
        placeholder_names = [
            "TITLE",
            "BODY",
            "CENTER_TITLE",
            "SUBTITLE",
            "DATE",
            "SLIDE_NUMBER",
            "FOOTER",
            "OBJECT",
            "CHART",
            "TABLE",
            "PICTURE",
            "CLIP_ART",
            "SMART_ART",
            "MEDIA",
        ]
        type_mapping = {}
        for name in placeholder_names:
            if hasattr(PP_PLACEHOLDER, name):
                type_mapping[getattr(PP_PLACEHOLDER, name)] = name

        return type_mapping.get(ph_type, "UNKNOWN")

    def get_layout_names(self) -> List[str]:
        """Get list of available layout names."""
        if not self.metadata:
            self.parse()
        return [layout["name"] for layout in self.metadata["slide_layouts"]]

    def get_placeholders_by_layout(self, layout_name: str) -> List[Dict[str, Any]]:
        """
        Get placeholders for a specific layout by name.

        Args:
            layout_name: Name of the slide layout

        Returns:
            List of placeholder dictionaries for that layout

        Raises:
            ValueError: If layout_name not found
        """
        if not self.metadata:
            self.parse()

        for layout in self.metadata["slide_layouts"]:
            if layout["name"] == layout_name:
                return layout["placeholders"]

        raise ValueError(f"Layout '{layout_name}' not found")
